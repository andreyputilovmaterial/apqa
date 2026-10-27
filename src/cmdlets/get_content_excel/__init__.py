"""`get-content-excel` — read an Excel workbook into a block of sheets.

Each sheet is keyed by its name and holds that sheet's data rows as an array of
blocks (one dict per data row, keyed by header cell). Returning every sheet —
keyed by its stable sheet name — rather than a single active sheet means a QC
rule must address the sheet it wants explicitly and never depends on which sheet
Excel happened to save as "active".

Optional dependency: ``openpyxl`` (see ``requirements.txt``). The import is
guarded so a missing dependency never breaks loading this module or listing the
cmdlet — only *running* the cmdlet raises a readable error (chained to the
original import error).
"""

from ..registry import cmdlets
from ..value_types import Block, PipeValue

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except Exception as error:  # noqa: BLE001 - any missing-dependency failure
    openpyxl = None
    get_column_letter = None
    _openpyxl_import_error = error
else:
    _openpyxl_import_error = None


@cmdlets.register(
    'get-content-excel',
    kind='source',
    help='Read every sheet of an Excel workbook into a block keyed by sheet name.',
    args=[
        {'name': 'path', 'type': 'path', 'required': True, 'description': 'Full path to the Excel workbook.'},
    ],
)
def get_content_excel(value: PipeValue, args: dict) -> Block:
    if _openpyxl_import_error is not None:
        raise ImportError(
            "get-content-excel: missing optional dependency 'openpyxl'"
        ) from _openpyxl_import_error

    path = args.get('path')
    if not path:
        raise ValueError("get-content-excel: missing required arg 'path'")

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            sheet_name: _read_sheet(workbook[sheet_name])
            for sheet_name in workbook.sheetnames
        }
    finally:
        close = getattr(workbook, 'close', None)
        if callable(close):
            close()


def _read_sheet(sheet) -> list:
    """Read one sheet into an array of row blocks (header row = field names)."""
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    names = _deduplicate_headers(header)
    return [dict(zip(names, row)) for row in rows]


def _deduplicate_headers(header) -> list:
    """Turn raw header cells into unique field names.

    Empty cells fall back to the source's native coordinate (Excel letters);
    duplicate names get a numeric suffix (``state``, ``state.1``, ``state.2``)
    per ``docs/cmdlet-data-formats.md``.
    """
    seen = {}
    names = []
    for index, cell in enumerate(header, start=1):
        base = cell if cell not in (None, '') else get_column_letter(index)
        base = f'{base}'
        if base in seen:
            seen[base] += 1
            names.append(f'{base}.{seen[base]}')
        else:
            seen[base] = 0
            names.append(base)
    return names
